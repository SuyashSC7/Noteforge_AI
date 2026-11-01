import sys
import os
from docx import Document
from pptx import Presentation

def read_docx(path):
    doc = Document(path)
    for paragraph in doc.paragraphs:
        print(paragraph.text)

    for table in doc.tables:
        for row in table.rows:
            sentence=[cell.text for cell in row.cells]
            print(" | ".join(sentence))


def read_pptx(path):

    prs=Presentation(path)

    for i , slide in enumerate( prs.slides , start=1 ):
        print(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if (shape.has_text_frame):
                print(shape.text_frame.text)
            
            if (shape.has_table):
                for row in shape.table.rows:
                    sentence=[cell.text for cell in row.cells]
                    print(" | ".join(sentence))


def main():
    if len(sys.argv)<2 :
        sys.exit(1)
    
    path = sys.argv[1]

    ext= os.path.splitext(path)[1].lower()
    try:
        if ext==".pptx":
            read_pptx(path)
        
        elif (ext==".docx"):
            read_docx(path)

        else:
            print("unsupported data type ", file=sys.stderr)
            sys.exit(1)
    except Exception as e:
        print(f"could not read the {path}: {e}", file=sys.stderr)
        sys.exit(1)


main()